"""Capital en existencia: valorización del inventario y planilla Excel.

Criterio (decidido con el usuario):
* Moneda base: USD. Cada precio se convierte con la cotización cargada en la tabla
  `cotizaciones` (unidades de la moneda por 1 USD). Nunca se inventa una cotización:
  si falta, el componente queda «Sin cotización» y no suma.
* Precio de cada componente: el local si existe (> 0); si no, el del exterior.
* Capital de un componente = precio unitario en USD × stock.

La planilla se escribe con FÓRMULAS, así que al cambiar una cotización dentro de Excel
se recalcula todo. `valorizar` repite en Python la misma regla para poder mostrar un
resumen antes de generar el archivo y para verificar que las fórmulas son correctas.

Este módulo no usa la interfaz. `openpyxl` se importa recién al exportar, así la
aplicación arranca aunque no esté instalado.
"""
from datetime import datetime

import db

SIN_DEFINIR = "(sin definir)"
ESTADO_OK = "OK"
ESTADO_SIN_PRECIO = "Sin precio"
ESTADO_SIN_COTIZACION = "Sin cotización"

INSTRUCCIONES_PIP = "Instale la biblioteca con:  pip install openpyxl"


class ExportError(Exception):
    """No se pudo generar la planilla; el mensaje explica el motivo."""


# -------------------------------------------------------------------
# Valorización
# -------------------------------------------------------------------
def _numero(valor):
    try:
        return float(valor)
    except (TypeError, ValueError):
        return 0.0


def valorizar(componente, cotizaciones):
    """Valora un componente. `cotizaciones`: {moneda en minúscula: unidades por 1 USD} (solo las cargadas).

    Devuelve un diccionario con: precio, moneda, origen ('Local'/'Exterior'/'—'), unitario_usd (o None),
    total_usd y estado.
    """
    local, exterior = _numero(componente.get("precio_local")), _numero(componente.get("precio_exterior"))
    if local > 0:
        precio, moneda, origen = local, str(componente.get("moneda_local") or "").strip(), "Local"
    elif exterior > 0:
        precio, moneda, origen = exterior, str(componente.get("moneda_exterior") or "").strip(), "Exterior"
    else:
        return dict(precio=0.0, moneda="", origen="—", unitario_usd=None, total_usd=0.0, estado=ESTADO_SIN_PRECIO)

    tasa = cotizaciones.get(moneda.lower())
    if not tasa or tasa <= 0:
        return dict(precio=precio, moneda=moneda, origen=origen, unitario_usd=None, total_usd=0.0, estado=ESTADO_SIN_COTIZACION)
    unitario = precio / tasa
    stock = int(_numero(componente.get("stock")))
    return dict(precio=precio, moneda=moneda, origen=origen, unitario_usd=unitario, total_usd=unitario * stock, estado=ESTADO_OK)


def calcular():
    """Valoriza todo el inventario con las cotizaciones cargadas.

    Devuelve: filas (componente + 'valor'), cotizaciones [(moneda, por_usd|None, fecha|None)], total_usd,
    unidades, y los conteos sin_precio / sin_cotizacion, con la lista de monedas sin cotización (que se usan).
    """
    cotizaciones = db.listar_cotizaciones()
    tasas = {m.lower(): p for m, p, _ in cotizaciones if p}
    filas = []
    for comp in db.listar_inventario():
        fila = dict(comp)
        fila["valor"] = valorizar(comp, tasas)
        filas.append(fila)
    sin_cotizacion = [f for f in filas if f["valor"]["estado"] == ESTADO_SIN_COTIZACION]
    return dict(
        filas=filas,
        cotizaciones=cotizaciones,
        total_usd=sum(f["valor"]["total_usd"] for f in filas),
        unidades=sum(int(_numero(f.get("stock"))) for f in filas),
        sin_precio=sum(1 for f in filas if f["valor"]["estado"] == ESTADO_SIN_PRECIO),
        sin_cotizacion=len(sin_cotizacion),
        monedas_sin_cotizacion=sorted({(f["valor"]["moneda"] or "(sin moneda)") for f in sin_cotizacion}),
    )


def formato_monto(valor):
    """1234567.5 -> '1.234.567,50' (como se escribe en Uruguay)."""
    return f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


# -------------------------------------------------------------------
# Planilla Excel
# -------------------------------------------------------------------
def _grupos(filas):
    """Tipos, categorías y subcategorías que realmente tienen componentes, ordenados sin distinguir mayúsculas."""
    def nivel(fila, i):
        return (fila[("tipo_componente", "categoria", "subcategoria")[i]] or "").strip() or SIN_DEFINIR
    caminos = {tuple(nivel(f, i) for i in range(k)) for f in filas for k in (1, 2, 3)}
    return sorted(caminos, key=lambda c: tuple(x.lower() for x in c))


def subtotales(filas):
    """Subtotales por clasificación (para mostrar en pantalla).

    Devuelve [(camino, total_usd, componentes, unidades)] con los mismos grupos de `_grupos`
    (padres antes que hijos). Solo suma los componentes valorizados.
    """
    def nivel(fila, i):
        return (fila[("tipo_componente", "categoria", "subcategoria")[i]] or "").strip() or SIN_DEFINIR
    resultado = []
    for camino in _grupos(filas):
        k = len(camino)
        miembros = [f for f in filas if tuple(nivel(f, i) for i in range(k)) == camino]
        resultado.append((camino, sum(f["valor"]["total_usd"] for f in miembros), len(miembros),
                          sum(int(_numero(f.get("stock"))) for f in miembros)))
    return resultado


def exportar_xlsx(destino, datos=None):
    """Escribe la planilla en `destino` (.xlsx). `datos`: resultado de calcular() (si no se da, se calcula).

    Hojas: Resumen (total y subtotales por clasificación), Detalle (un renglón por componente) y
    Cotizaciones (las que se usaron; editables). Todo lo calculado son fórmulas.
    Lanza ExportError si falta openpyxl; los errores de escritura (archivo abierto en Excel, sin
    permiso) salen como OSError.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.properties import PageSetupProperties
    except ImportError:
        raise ExportError("Falta la biblioteca openpyxl, necesaria para generar la planilla.\n" + INSTRUCCIONES_PIP)

    datos = datos or calcular()
    filas, cotizaciones = datos["filas"], datos["cotizaciones"]
    otras = [m for m, _, _ in cotizaciones if m.upper() != db.MONEDA_BASE]

    # --- estilos ---
    arial = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
    f_normal, f_negrita = arial(), arial(bold=True)
    f_titulo, f_cab = arial(size=14, bold=True), arial(bold=True, color="FFFFFF")
    f_entrada = arial(color="0000FF")              # valores que el usuario puede cambiar
    f_gris = arial(color="808080")
    fill_cab = PatternFill("solid", start_color="1F3864")
    fill_tipo = PatternFill("solid", start_color="D9E1F2")
    fill_cat = PatternFill("solid", start_color="EEF2F9")
    fill_falta = PatternFill("solid", start_color="FFFF00")
    borde = Border(bottom=Side(style="thin", color="BFBFBF"))
    f_monto, f_pct, f_ent = "#,##0.00", "0.0%", "#,##0"

    wb = Workbook()
    ws_res, ws_det = wb.active, wb.create_sheet("Detalle")
    ws_res.title = "Resumen"
    ws_cot = wb.create_sheet("Cotizaciones")

    # ================= Cotizaciones =================
    ws_cot["A1"], ws_cot["A1"].font = "Cotizaciones", f_titulo
    ws_cot["A2"] = ("Unidades de cada moneda que equivalen a 1 USD (ej.: UYU 40,50 significa 1 USD = 40,50 UYU). "
                    "Los valores en azul se pueden cambiar: toda la planilla se recalcula. En amarillo, las que faltan cargar.")
    ws_cot["A2"].font = f_gris
    for c, titulo in enumerate(("Moneda", "Unidades por 1 USD", "Actualizada"), start=1):
        celda = ws_cot.cell(row=4, column=c, value=titulo)
        celda.font, celda.fill = f_cab, fill_cab
    for i, (moneda, por_usd, fecha) in enumerate(cotizaciones):
        r = 5 + i
        ws_cot.cell(row=r, column=1, value=moneda).font = f_normal
        celda = ws_cot.cell(row=r, column=2, value=por_usd)
        celda.number_format = "#,##0.0000"
        if moneda.upper() == db.MONEDA_BASE:
            celda.font = f_normal
        else:
            celda.font = f_entrada
            if por_usd is None:
                celda.fill = fill_falta
        ws_cot.cell(row=r, column=3, value=fecha or "").font = f_normal
    COT = "Cotizaciones!$A$5:$B$60"                 # holgura para monedas que se agreguen dentro de Excel
    ws_cot.column_dimensions["A"].width, ws_cot.column_dimensions["B"].width, ws_cot.column_dimensions["C"].width = 14, 22, 16

    # ================= Detalle =================
    cab = ["ID", "Tipo", "Categoría", "Subcategoría", "Ítem", "Chip Principal", "Descripción", "Fabricante", "Stock",
           "Precio Local", "Moneda Local", "Precio Exterior", "Moneda Exterior",
           "Precio usado", "Moneda usada", "Origen", "Precio unit. (USD)", "Total (USD)"]
    cab += [f"Total ({m})" for m in otras] + ["Estado"]
    for c, titulo in enumerate(cab, start=1):
        celda = ws_det.cell(row=1, column=c, value=titulo)
        celda.font, celda.fill = f_cab, fill_cab
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws_det.row_dimensions[1].height = 32

    col_estado = len(cab)
    letra_estado = get_column_letter(col_estado)
    for i, comp in enumerate(filas):
        r = 2 + i
        nivel = lambda clave: (comp[clave] or "").strip() or SIN_DEFINIR
        valores = [comp["id"], nivel("tipo_componente"), nivel("categoria"), nivel("subcategoria"), (comp["item_especifico"] or ""),
                   comp["chip_principal"] or "", comp["descripcion"] or "", comp["fabricante"] or "",
                   int(_numero(comp["stock"])),
                   comp["precio_local"] or 0, comp["moneda_local"] or "", comp["precio_exterior"] or 0, comp["moneda_exterior"] or ""]
        for c, v in enumerate(valores, start=1):
            ws_det.cell(row=r, column=c, value=v).font = f_normal
        formulas = [
            f'=IF(J{r}>0,J{r},IF(L{r}>0,L{r},0))',                                              # N precio usado
            f'=IF(J{r}>0,K{r},IF(L{r}>0,M{r},""))',                                             # O moneda usada
            f'=IF(J{r}>0,"Local",IF(L{r}>0,"Exterior","—"))',                                   # P origen
            f'=IF(N{r}>0,IFERROR(N{r}/VLOOKUP(O{r},{COT},2,FALSE),""),"")',                      # Q unitario USD
            f'=IF(ISNUMBER(Q{r}),Q{r}*I{r},0)',                                                 # R total USD
        ]
        for m in otras:
            formulas.append(f'=IFERROR(IF(AND(ISNUMBER(Q{r}),VLOOKUP("{m}",{COT},2,FALSE)>0),R{r}*VLOOKUP("{m}",{COT},2,FALSE),""),"")')
        formulas.append(f'=IF(N{r}<=0,"{ESTADO_SIN_PRECIO}",IF(ISNUMBER(Q{r}),"{ESTADO_OK}","{ESTADO_SIN_COTIZACION}"))')
        for c, formula in enumerate(formulas, start=14):
            ws_det.cell(row=r, column=c, value=formula).font = f_normal
        for c in (10, 12, 14, 17, 18) + tuple(range(19, col_estado)):
            ws_det.cell(row=r, column=c).number_format = f_monto
        ws_det.cell(row=r, column=9).number_format = f_ent
    n = len(filas)
    ultimo = max(n + 1, 2)
    # franjas alternadas (formato condicional: se mantienen al ordenar o filtrar)
    from openpyxl.formatting.rule import FormulaRule
    rango_franjas = f"A2:{letra_estado}{ultimo}"
    ws_det.conditional_formatting.add(rango_franjas, FormulaRule(formula=["MOD(ROW(),2)=0"], fill=PatternFill("solid", start_color="FFF9D9", end_color="FFF9D9")))
    ws_det.conditional_formatting.add(rango_franjas, FormulaRule(formula=["MOD(ROW(),2)=1"], fill=PatternFill("solid", start_color="E3ECF2", end_color="E3ECF2")))
    for c, ancho in enumerate([6, 24, 28, 28, 26, 18, 34, 22, 8, 12, 9, 12, 9, 12, 9, 10, 14, 14], start=1):
        ws_det.column_dimensions[get_column_letter(c)].width = ancho
    for c in range(19, col_estado + 1):
        ws_det.column_dimensions[get_column_letter(c)].width = 15
    ws_det.freeze_panes = "C2"
    ws_det.auto_filter.ref = f"A1:{letra_estado}{ultimo}"

    rango = lambda letra: f"Detalle!${letra}$2:${letra}${ultimo}"

    # ================= Resumen =================
    ws_res["A1"], ws_res["A1"].font = "Capital en existencia", f_titulo
    ws_res["A2"] = (f"Generado el {datetime.now():%d/%m/%Y %H:%M}  ·  Moneda base: {db.MONEDA_BASE}  ·  "
                    "Precio usado: el local si existe, si no el del exterior  ·  Capital = precio unitario × stock")
    ws_res["A2"].font = f_gris

    kpis = [
        ("Capital total (USD)", f"=SUM({rango('R')})", f_monto),
    ]
    for m in otras:
        kpis.append((f"Equivalente en {m}", f'=IFERROR(IF(VLOOKUP("{m}",{COT},2,FALSE)>0,$B$4*VLOOKUP("{m}",{COT},2,FALSE),"sin cotización"),"sin cotización")', f_monto))
    kpis += [
        ("Componentes", f"=COUNTA({rango('A')})", f_ent),
        ("Unidades en existencia", f"=SUM({rango('I')})", f_ent),
        ("Sin precio (no suman)", f'=COUNTIF({rango(letra_estado)},"{ESTADO_SIN_PRECIO}")', f_ent),
        ("Sin cotización (no suman)", f'=COUNTIF({rango(letra_estado)},"{ESTADO_SIN_COTIZACION}")', f_ent),
    ]
    for i, (rotulo, formula, formato) in enumerate(kpis):
        r = 4 + i
        ws_res.cell(row=r, column=1, value=rotulo).font = f_negrita if i == 0 else f_normal
        celda = ws_res.cell(row=r, column=2, value=formula)
        celda.number_format, celda.font = formato, (arial(bold=True, size=12) if i == 0 else f_normal)
        celda.alignment = Alignment(horizontal="right")

    r0 = 4 + len(kpis) + 1
    ws_res.cell(row=r0, column=1, value="Subtotales por clasificación").font = arial(bold=True, size=12)
    r_cab = r0 + 1
    for c, titulo in enumerate(("Tipo", "Categoría", "Subcategoría", "Componentes", "Unidades", "Capital (USD)", "% del total"), start=1):
        celda = ws_res.cell(row=r_cab, column=c, value=titulo)
        celda.font, celda.fill = f_cab, fill_cab
        celda.alignment = Alignment(horizontal="center")

    letras_nivel = ("B", "C", "D")      # columnas de Detalle: Tipo, Categoría, Subcategoría
    r = r_cab + 1
    filas_tipo = []
    for camino in _grupos(filas):
        k = len(camino)
        for i, nombre in enumerate(camino):
            celda = ws_res.cell(row=r, column=1 + i, value=nombre)
            celda.font = f_negrita if i == k - 1 else f_gris
        cond = "*".join(f"({rango(letras_nivel[i])}=${get_column_letter(1 + i)}{r})" for i in range(k))
        ws_res.cell(row=r, column=4, value=f"=SUMPRODUCT({cond}*1)")
        ws_res.cell(row=r, column=5, value=f"=SUMPRODUCT({cond}*{rango('I')})")
        ws_res.cell(row=r, column=6, value=f"=SUMPRODUCT({cond}*{rango('R')})")
        ws_res.cell(row=r, column=7, value=f"=IF($B$4=0,0,F{r}/$B$4)")
        for c, formato in ((4, f_ent), (5, f_ent), (6, f_monto), (7, f_pct)):
            ws_res.cell(row=r, column=c).number_format = formato
            ws_res.cell(row=r, column=c).font = f_negrita if k == 1 else f_normal
        relleno = fill_tipo if k == 1 else (fill_cat if k == 2 else None)
        for c in range(1, 8):
            ws_res.cell(row=r, column=c).border = borde
            if relleno:
                ws_res.cell(row=r, column=c).fill = relleno
        if k == 1:
            filas_tipo.append(r)
        r += 1

    # control: los subtotales por Tipo deben sumar el total
    r += 1
    suma = "+".join(f"F{x}" for x in filas_tipo) or "0"
    ws_res.cell(row=r, column=1, value="Control: suma de los Tipos").font = f_normal
    ws_res.cell(row=r, column=6, value=f"={suma}").number_format = f_monto
    ws_res.cell(row=r, column=7, value=f'=IF(ABS(F{r}-$B$4)<0.005,"OK","Revisar")').font = f_negrita
    ws_res.cell(row=r, column=6).font = f_normal

    for letra, ancho in zip("ABCDEFG", (32, 52, 34, 14, 12, 18, 12)):
        ws_res.column_dimensions[letra].width = ancho
    ws_res.freeze_panes = ws_res.cell(row=r_cab + 1, column=1)

    for hoja in (ws_res, ws_det, ws_cot):      # al imprimir: horizontal y todas las columnas en el ancho de la hoja
        hoja.page_setup.orientation = "landscape"
        hoja.page_setup.fitToWidth, hoja.page_setup.fitToHeight = 1, 0
        hoja.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws_det.print_title_rows = "1:1"

    wb.calculation.fullCalcOnLoad = True       # Excel calcula las fórmulas al abrir
    wb.save(destino)
    return destino
